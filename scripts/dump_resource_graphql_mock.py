#!/usr/bin/env python3
"""Dump the file-based Resource service as a GetNightConfigurations response.

Read-only. It loads OcsResourceService / SimResourceService / OpsResourceService,
re-reads the telescope schedule spreadsheet for the raw 'Mode/Program' column
(the service consumes that column into filters and does not keep it), and writes
JSON shaped exactly like the query in
backend/scheduler/services/resource/graphql/queries.graphql.

The point is to make the endpoint swap cheap and *checkable*: a GraphQL-backed
ResourceService can be developed against this fixture, and once GPP serves the
real thing the two payloads can be diffed night by night.

Two places where the fixture is *not* the raw observatory record, because it is
read back out of the file service rather than out of the files:

  * ToO triggers that fired before evening twilight already carry the clamped
    time (eve_twi + 25s), not the real trigger instant. Harmless for scheduling
    (clamping again is a no-op) but the real endpoint should serve the raw time.
  * Faults and weather closures shorter than 20 minutes are already gone. The
    real endpoint should serve them and let the Scheduler drop them.

Usage (from the repo root):

    .venv/bin/python scripts/dump_resource_graphql_mock.py --mode operation
    .venv/bin/python scripts/dump_resource_graphql_mock.py --mode validation \
        --site GN --start 2018-08-01 --end 2018-08-31 -o /tmp/gn_aug.json
"""

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "backend"))

from lucupy.helpers import remove_symbols, str_to_bool  # noqa: E402
from lucupy.minimodel import ProgramID, Site, TimeAccountingCode  # noqa: E402
from openpyxl.reader.excel import load_workbook  # noqa: E402

from definitions import ROOT_DIR  # noqa: E402
from scheduler.services.resource import (
    OcsResourceService,
    OpsResourceService,  # noqa: E402
    SimResourceService,
)

_SERVICES = {
    "validation": OcsResourceService,
    "operation": OpsResourceService,
    "simulation": SimResourceService,
}

# ITCD name -> long/OCS name, inverted from the tables on ResourceService so the
# fpuBarcodes query can carry both spellings.
_LONG_NAMES = {
    Site.GN: {
        v: k for k, v in dict(getattr(OcsResourceService, "_gmosn_ifu_dict")).items()
    },
    Site.GS: {
        v: k for k, v in dict(getattr(OcsResourceService, "_gmoss_ifu_dict")).items()
    },
}


def _resource(r) -> Dict:
    return {"id": r.id, "description": r.description, "type": r.type.name}


def _interruption(i, fault_report: Optional[str] = None) -> Dict:
    return {
        "start": i.start_time.isoformat(),
        "end": i.end_time.isoformat(),
        "description": i.description,
        "faultReport": fault_report,
        "affects": sorted((_resource(r) for r in i.affects), key=lambda d: d["id"]),
    }


def _fault(f) -> Dict:
    # The file loader folds the FR number into the description as 'FR-1234: text'.
    # The endpoint carries it as its own field, so split it back out.
    match = re.match(r"(FR-\d+):\s*(.*)", f.description or "")
    if match:
        return _interruption(f, fault_report=match.group(1)) | {
            "description": match.group(2)
        }
    return _interruption(f)


def _read_schedule(subdir: str, site: Site) -> Dict[date, Dict]:
    """Re-parse the 'Telescope' and 'Mode/Program' columns, which the service does not retain."""
    path = (
        Path(ROOT_DIR)
        / "scheduler"
        / "services"
        / "resource"
        / "data"
        / subdir
        / "telescope_schedules.xlsx"
    )
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    sheet = workbook[site.name]

    schedule: Dict[date, Dict] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        row_date = row[0].date()
        status = str(row[1]).upper().strip()
        entry_list: List[Dict] = []

        if status == "CLOSED":
            schedule[row_date] = {"telescopeStatus": "CLOSED", "modes": entry_list}
            continue

        # Parenthetical expressions are dropped by the service too.
        modes_entry = re.sub(r"\([^)]*\)", "", row[2]).upper().strip()
        for entry in (e.strip() for e in modes_entry.split("|")):
            if entry.startswith("VISITOR:"):
                entry_list.append(
                    {
                        "type": "VISITOR",
                        "programs": [],
                        "partner": None,
                        "instrument": {
                            "id": remove_symbols(entry[8:].strip()),
                            "description": None,
                            "type": "INSTRUMENT",
                        },
                    }
                )
            elif entry.startswith("PARTNER:"):
                entry_list.append(
                    {
                        "type": "PARTNER",
                        "programs": [],
                        "partner": TimeAccountingCode[entry[8:].strip()].name,
                        "instrument": None,
                    }
                )
            elif entry.startswith("PV:"):
                entry_list.append(_program_mode("PV", entry[3:]))
            elif entry.startswith("CLASSICAL:"):
                entry_list.append(_program_mode("CLASSICAL", entry[10:]))
            elif entry.startswith("PRIORITY:"):
                entry_list.append(_program_mode("PRIORITY", entry[9:]))
            elif entry in ("ENGINEERING", "SHUTDOWN", "QUEUE"):
                entry_list.append(
                    {"type": entry, "programs": [], "partner": None, "instrument": None}
                )

        schedule[row_date] = {"telescopeStatus": "OPEN", "modes": entry_list}
    return schedule


def _program_mode(mode_type: str, text: str) -> Dict:
    programs = sorted(ProgramID(p.strip()).id for p in text.strip().split(","))
    return {
        "type": mode_type,
        "programs": programs,
        "partner": None,
        "instrument": None,
    }


def dump(
    subdir: str, sites: List[Site], start: Optional[date], end: Optional[date]
) -> Dict:
    service = _SERVICES[subdir](frozenset(sites))
    payload = {"data": {}}

    for site in sites:
        schedule = _read_schedule(subdir, site)
        earliest, latest = service.date_range_for_site(site)
        first = max(start, earliest) if start else earliest
        last = min(end, latest) if end else latest

        matches = []
        d = first
        while d <= last:
            blocked = d in service._blocked[site]
            night = schedule.get(d, {"telescopeStatus": "OPEN", "modes": []})
            matches.append(
                {
                    "site": site.name,
                    "localDate": d.isoformat(),
                    "telescopeStatus": night["telescopeStatus"],
                    "blocked": blocked,
                    "laserGuideStar": (not blocked)
                    and service._lgs[site].get(d, False),
                    "tooAccepted": (not blocked) and service._too[site].get(d, True),
                    "modes": night["modes"],
                    "resources": sorted(
                        (_resource(r) for r in service.get_resources(site, d)),
                        key=lambda x: (x["type"], x["id"]),
                    ),
                    "engineeringTasks": sorted(
                        (
                            {
                                "start": t.start_time.isoformat(),
                                "end": t.end_time.isoformat(),
                                "description": t.description,
                            }
                            for t in service.get_eng_tasks(site, d)
                        ),
                        key=lambda x: x["start"],
                    ),
                    "faults": sorted(
                        (_fault(f) for f in service.get_faults(site, d)),
                        key=lambda x: x["start"],
                    ),
                    "weatherClosures": sorted(
                        (
                            _interruption(w)
                            for w in service.get_unexpected_closures(site, d)
                        ),
                        key=lambda x: x["start"],
                    ),
                    "tooActivations": sorted(
                        (
                            {
                                "observationId": t.too_id.id,
                                "triggeredAt": t.start_time.isoformat(),
                            }
                            for t in service.get_toos(site, d)
                        ),
                        key=lambda x: x["triggeredAt"],
                    ),
                }
            )
            d += type(service)._day

        payload["data"][f"nightConfigurations_{site.name}"] = {
            "matches": matches,
            "hasMore": False,
        }
        payload["data"][f"fpuBarcodes_{site.name}"] = [
            {
                "itcdName": itcd,
                "barcode": res.id,
                "longName": _LONG_NAMES[site].get(itcd),
            }
            for itcd, res in sorted(service._itcd_fpu_to_barcode[site].items())
        ]

    return payload


def main() -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--mode", choices=sorted(_SERVICES), default="operation")
    parser.add_argument(
        "--site",
        choices=["GN", "GS"],
        action="append",
        help="Repeatable. Defaults to both sites.",
    )
    parser.add_argument(
        "--start", type=lambda s: datetime.strptime(s, "%Y-%m-%d").date()
    )
    parser.add_argument("--end", type=lambda s: datetime.strptime(s, "%Y-%m-%d").date())
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="Defaults to <mode>_night_configurations.json in the CWD.",
    )
    args = parser.parse_args()

    sites = [Site[s] for s in (args.site or ["GN", "GS"])]
    payload = dump(args.mode, sites, args.start, args.end)

    output = args.output or Path(f"{args.mode}_night_configurations.json")
    output.write_text(json.dumps(payload, indent=2))

    nights = sum(
        len(v["matches"])
        for k, v in payload["data"].items()
        if k.startswith("nightConfigurations")
    )
    print(f"Wrote {nights} nights for {[s.name for s in sites]} to {output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
